import copy
import json
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from cashwatch.data import audit_sequence_pair, file_hash, group_review_status, prepare, read_annotations, write_json
from tools.reserve_test import choose_test_groups


class ExpansionTests(unittest.TestCase):
    def test_csv_requires_explicit_owner_confirmation_and_carry_is_limited(self):
        record = {"id": "a", "group": "g", "sha256": "h"}
        self.assertEqual(group_review_status([record], {"a"}, None, [], True)["status"], "pending")
        confirmed = group_review_status([record], {"a"}, None, [], True, True)
        self.assertEqual(confirmed["status"], "confirmed")
        old = {"records": [record], "group_review": confirmed}
        self.assertEqual(group_review_status([record], {"a"}, old, [])["status"], "confirmed")
        self.assertEqual(group_review_status([record], {"a", "new"}, old, [])["status"], "pending")
        self.assertEqual(group_review_status([record], {"a"}, old, [{"id": "a"}])["status"], "pending")
        legacy = {"records": [record], "group_review": {"status": "confirmed"}}
        self.assertEqual(group_review_status([record], {"a"}, legacy, [])["status"], "pending")
        with self.assertRaisesRegex(ValueError, "requires --groups-csv"):
            prepare("not_read", "not_created", group_review_confirmed=True)

    def test_locked_label_revision_is_explicit_archived_and_invalidates_test(self):
        with tempfile.TemporaryDirectory() as folder:
            root, output = Path(folder), Path(folder) / "output"
            records = []
            for name, split, label, times in [("x-1", "train", 0, []), ("x-2", "dev", 1, [1.0]), ("x-3", "test", 0, [])]:
                path = root / f"{name}.mov"
                path.write_bytes(name.encode())
                digest = file_hash(path)
                records.append({"id": name, "sha256": digest, "group": digest, "split": split,
                                "label": label, "event_times": times})
            old = {"seed": 42, "records": records, "labels_sha256": "old-label-hash",
                   "test_reservation": {"selected_ids": ["x-3"], "status": "provenance_confirmed"}}
            write_json(output / "manifest.json", old)
            initial_bytes = (output / "manifest.json").read_bytes()
            workbook = Workbook()
            workbook.active.append(["наименование файла", "таймкод", "метка"])
            workbook.active.append(["x-1", None, "без нарушения"])
            workbook.active.append(["x-2", 1.5, "нарушение"])
            workbook.active.append(["x-3", 2, "нарушение"])
            workbook.save(root / "labels.xlsx")
            workbook.close()
            groups = root / "groups.csv"
            groups.write_text("filename,group\nx-1,one\nx-2,two\nx-3,three\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "--accept-label-revisions"):
                prepare(root, output)
            self.assertEqual((output / "manifest.json").read_bytes(), initial_bytes)
            images = np.zeros((2, 224, 224, 3), dtype=np.uint8)
            with patch("cashwatch.data.decode_sampled", return_value=(images, np.array([0., 0.5]), 4.)), \
                    patch("cashwatch.data.audit_sequences", return_value={"confirmed_overlap_pairs": []}), \
                    patch("cashwatch.data.resolve_groups", return_value=[]):
                updated = prepare(root, output, groups_csv=groups, group_review_confirmed=True, accept_label_revisions=True)
            self.assertEqual({r["id"]: r["split"] for r in updated["records"]}, {r["id"]: r["split"] for r in records})
            self.assertEqual(updated["group_review"]["status"], "pending")
            self.assertEqual(updated["test_reservation"]["qualification_status"], "invalidated_by_label_revision")
            self.assertEqual(updated["test_reservation"]["status"], "provenance_pending")
            saved_revision = Path(updated["label_revision_history"][0]["previous_manifest"])
            self.assertEqual(json.loads(saved_revision.read_text(encoding="utf-8")), old)
            self.assertEqual(len(updated["annotation_changes_since_previous_manifest"]), 2)

    def test_prepare_refuses_concurrently_changed_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook = Workbook()
            workbook.active.append(["наименование файла", "таймкод", "метка"])
            workbook.active.append(["X-1", None, "без нарушения"])
            workbook.save(root / "labels.xlsx")
            workbook.close()
            with patch("cashwatch.data.file_hash", side_effect=["before", "after"]):
                with self.assertRaisesRegex(ValueError, "changed while being read"):
                    prepare(root, root / "output")
            self.assertFalse((root / "output" / "manifest.json").exists())

    def test_temporal_crop_reencoding_and_shuffled_frames(self):
        rng = np.random.default_rng(7)
        source = rng.uniform(20, 220, (30, 8, 8, 3)).astype(np.float32)
        cropped = source[8:23] + rng.normal(0, 0.3, (15, 8, 8, 3))
        result = audit_sequence_pair(source, cropped, np.arange(30) / 2, np.arange(15) / 2)
        self.assertEqual(result["status"], "confirmed_numeric_overlap")
        self.assertEqual(result["matched_frames"], 15)
        self.assertEqual(result["offset_sec"], -4)
        self.assertIsNone(audit_sequence_pair(source, cropped[::-1], np.arange(30) / 2, np.arange(15) / 2))

    def test_short_sequences_do_not_match_or_index_past_end(self):
        long = np.zeros((30, 8, 8, 3), dtype=np.float32)
        for length in (0, 1, 7):
            short = long[:length]
            self.assertIsNone(audit_sequence_pair(short, long, np.arange(length) / 2, np.arange(30) / 2))
            self.assertIsNone(audit_sequence_pair(long, short, np.arange(30) / 2, np.arange(length) / 2))

    def test_same_static_background_does_not_prove_overlap(self):
        source = np.full((20, 8, 8, 3), 100, dtype=np.float32)
        result = audit_sequence_pair(source, source, np.arange(20) / 2, np.arange(20) / 2)
        self.assertEqual(result["status"], "possible_numeric_overlap")
        self.assertEqual(result["temporal_change"], 0)

    def test_test_reservation_rejects_used_bytes_aliases_and_ambiguous_groups(self):
        records = [{"id": f"v{i}", "sha256": f"h{i}", "group": f"g{i}", "label": i % 2, "split": "train"}
                   for i in range(12)]
        historical = [dict(records[0])]
        records[1]["split"] = "dev"
        records[2]["sha256"] = historical[0]["sha256"]
        records[3]["group"] = historical[0]["group"]
        before = copy.deepcopy(records)
        selected, report = choose_test_groups(records, historical, ambiguous_ids=["v4"], per_class=2)
        self.assertEqual(records, before)
        self.assertTrue(set(selected).isdisjoint({"g0", "g1", "g2", "g4"}))
        self.assertEqual(selected, choose_test_groups(records, historical, ["v4"], per_class=2)[0])
        self.assertEqual(len(selected), 4)
        self.assertIn("unresolved_possible_numeric_overlap", report["excluded_groups"]["g4"])
        with self.assertRaises(ValueError):
            choose_test_groups(records, historical, [f"v{i}" for i in range(12)], per_class=2)
        records[-1]["split"] = "test"
        with self.assertRaises(ValueError):
            choose_test_groups(records, historical)

    def test_unclear_label_is_quarantined_not_negative_and_duplicate_rows_merge(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "labels.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["наименование файла", "таймкод", "метка"])
            sheet.append(["Х-1", "1 секунда", "неясно"])
            sheet.append(["X-2", "2", "нарушение"])
            sheet.append(["Х-2", "3", "нарушение"])
            workbook.save(path)
            workbook.close()
            rows = read_annotations(path)
            self.assertIsNone(rows[0]["label"])
            self.assertEqual(rows[0]["annotation_status"], "uncertain_quarantined")
            self.assertEqual(rows[1]["event_times"], [2, 3])
            self.assertEqual(len(rows), 2)


if __name__ == "__main__":
    unittest.main()
