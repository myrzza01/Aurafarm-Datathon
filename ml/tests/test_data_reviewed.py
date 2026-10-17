"""Reviewed workbook targets, provenance and deduplication must remain auditable."""
import copy
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook

from cashwatch.data import parse_timestamps, read_reviewed_annotations
from cashwatch.reviewed_data import assign_reviewed_split, prepare_reviewed, reconcile_exact_duplicates


HEADERS = ["Файл", "Исходная метка", "Исходный таймкод", "Группа дублей", "Статус проверки",
           "Метка после проверки", "Таймкод после проверки, с", "Тип события", "Проверяющий", "Комментарий", "Примечание к данным"]


def workbook_fixture(path, rows, snapshots=()):
    workbook = Workbook()
    review = workbook.active
    review.title = "Проверка видео"
    review.append(HEADERS)
    original = workbook.create_sheet("Исходная разметка")
    original.append(["наименование файла", "таймстемп", "нарушение / без нарушения"])
    for row in rows:
        review.append(row)
        original.append([row[0], row[2], row[1]])
    snapshot = workbook.create_sheet("Для ML")
    snapshot.append(["video_file", "sha256", "classification_ready", "label_id", "duplicate_files"])
    for row in snapshots:
        snapshot.append(row)
    workbook.save(path)
    workbook.close()


def record(identifier, label=1, times=(5.0,), type_id="safe_left_open", verified=False):
    return {"id": identifier, "label": label, "sha256": "same-bytes", "path": f"/{identifier}.mov",
        "label_verified": verified, "label_source": "human_confirmed" if verified else "source_original_weak",
        "review_status": "Подтверждено" if verified else "Спорно", "event_times": list(times) if label else [],
        "events": [{"timecode_sec": timestamp, "type_id": type_id, "type_verified": bool(type_id),
                    "annotation_provenance": {"row": identifier}} for timestamp in times] if label else [],
        "type_id": type_id, "type_verified": bool(type_id), "annotation_provenance": [{"row": identifier}],
        "original_annotation": {"label": label, "times": list(times)}}


class ReviewedAnnotationTests(unittest.TestCase):
    def test_natural_minutes_are_supported_without_accepting_ranges(self):
        self.assertEqual(parse_timestamps("1минута 8 секунда; 2 минуты 0 секунд"), [68, 120])
        self.assertEqual(parse_timestamps("1 час 2 минуты 3,5 секунды"), [3723.5])
        with self.assertRaises(ValueError):
            parse_timestamps("10-15 секунд")

    def test_confirmed_current_review_overrides_stale_snapshot_and_original(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            workbook_fixture(path, [["Х-1.mov", "нарушение", "22 секунда", None, "Подтверждено",
                "без нарушения", None, "без нарушения", "reviewer", "Human note", None]],
                [["X-1.mov", "old-sha", 0, 1, None]])
            before = path.read_bytes()
            imported = read_reviewed_annotations(path, allow_weak_original=True)
            target = imported["records"][0]
            self.assertEqual((target["id"], target["label"], target["event_times"]), ("x-1", 0, []))
            self.assertTrue(target["label_verified"])
            self.assertEqual(target["source_snapshot"]["label_id"], 1)
            self.assertEqual(target["original_annotation"]["label"], 1)
            self.assertEqual(target["annotation_provenance"][0]["sheet"], "Проверка видео")
            self.assertEqual(path.read_bytes(), before)

    def test_weak_original_and_action_confirmation_are_separate_explicit_choices(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            workbook_fixture(path, [["X-1.mov", "нарушение", 22, None, "Спорно", "без нарушения", None,
                "Не отдал полную сумму", "reviewer", "Confirmed human action note", None]])
            excluded = read_reviewed_annotations(path)["records"][0]
            self.assertIsNone(excluded["label"])
            weak = read_reviewed_annotations(path, allow_weak_original=True)["records"][0]
            self.assertEqual((weak["label"], weak["event_times"]), (1, [22]))
            self.assertFalse(weak["label_verified"])
            self.assertFalse(weak["type_verified"])
            final = read_reviewed_annotations(path, allow_weak_original=True, human_action_notes_confirmed=True)["records"][0]
            self.assertFalse(final["label_verified"])
            self.assertTrue(final["type_verified"])
            self.assertEqual(final["type_id"], "cash_not_fully_delivered")
            self.assertEqual(final["action_confirmation_source"], "explicit_user_instruction_in_current_conversation")

    def test_explicit_time_inside_confirmed_action_note_has_its_own_source(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            workbook_fixture(path, [["Х-47.mov", "нарушение", "1минута 8 секунда", None, "Спорно", "нарушение", None,
                "68 с — объединение банкнот разных клиентов", None, "Human action", None]])
            imported = read_reviewed_annotations(path, allow_weak_original=True, human_action_notes_confirmed=True)["records"][0]
            self.assertEqual(imported["event_times"], [68])
            self.assertEqual(imported["events"][0]["time_source"], "human_action_note_explicit_point")
            self.assertTrue(imported["events"][0]["time_verified"])


class ReviewedDuplicateTests(unittest.TestCase):
    def test_distinct_confirmed_actions_on_exact_duplicate_are_two_events_one_sample(self):
        first = record("x-4", times=(57,), type_id="accepted_cash_left_exposed", verified=True)
        alias = record("x-47", times=(68,), type_id="customer_cash_combined")
        before = copy.deepcopy([first, alias])
        kept, quarantined, audit = reconcile_exact_duplicates([first, alias])
        self.assertEqual([first, alias], before)
        self.assertEqual((len(kept), len(quarantined)), (1, 0))
        self.assertEqual(kept[0]["event_times"], [57, 68])
        self.assertEqual({e["type_id"] for e in kept[0]["events"]}, {"accepted_cash_left_exposed", "customer_cash_combined"})
        self.assertEqual(len(kept[0]["annotation_provenance"]), 2)
        self.assertEqual(audit[0]["sample_weight_count"], 1)
        self.assertEqual(audit[0]["reason"], "distinct_user_confirmed_actions_merged")

    def test_unresolved_same_action_time_and_binary_conflicts_are_quarantined(self):
        first = record("x-4", times=(57,))
        alias = record("x-47", times=(68,))
        kept, quarantined, audit = reconcile_exact_duplicates([first, alias])
        self.assertEqual(kept, [])
        self.assertEqual(quarantined[0]["exclusion_reason"], "unresolved_duplicate_event_time_conflict")
        alias["label"] = 0
        kept, quarantined, audit = reconcile_exact_duplicates([first, alias])
        self.assertEqual(kept, [])
        self.assertEqual(audit[0]["reason"], "unresolved_original_label_conflict")

    def test_confirmed_review_resolves_older_time_without_erasing_alias_provenance(self):
        first = record("x-4", times=(57,), type_id=None, verified=True)
        alias = record("x-47", times=(68,), type_id=None)
        kept, quarantined, audit = reconcile_exact_duplicates([first, alias])
        self.assertFalse(quarantined)
        self.assertEqual(kept[0]["event_times"], [57])
        self.assertEqual(kept[0]["alias_annotations"][1]["event_times"], [68])

    def test_conflicting_confirmed_binary_labels_cannot_be_silently_resolved(self):
        positive = record("x-4", verified=True)
        negative = record("x-47", label=0, verified=True)
        kept, quarantined, audit = reconcile_exact_duplicates([positive, negative])
        self.assertFalse(kept)
        self.assertEqual(quarantined[0]["exclusion_reason"], "conflicting_confirmed_labels")

    def test_numeric_and_non_numeric_aliases_sort_deterministically(self):
        kept, quarantined, _ = reconcile_exact_duplicates([record("camera"), record("x-47")])
        self.assertFalse(quarantined)
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0]["id"], "x-47")


class ReviewedPreparationTests(unittest.TestCase):
    def test_three_group_semantic_class_is_not_consumed_entirely_by_dev(self):
        records = []
        for kind, count in [("door", 3), ("cash", 5), ("rare_a", 1), ("rare_b", 1), (None, 6)]:
            for index in range(count):
                target = record(f"{kind}-{index}", label=int(kind is not None), type_id=kind)
                target["group"] = target["id"]
                records.append(target)
        assign_reviewed_split(records, holdout_fraction=0.25)
        for kind in ("door", "cash"):
            self.assertEqual({r["split"] for r in records if r["type_id"] == kind}, {"train", "dev", "test"})
        self.assertTrue(all(r["split"] == "train" for r in records if r["type_id"] in {"rare_a", "rare_b"}))

    def test_split_is_deterministic_disjoint_and_rare_types_stay_in_train(self):
        records = []
        for index in range(20):
            target = record(f"v{index}", label=index % 2, type_id="singleton" if index == 1 else "common")
            target["group"] = f"g{index}"
            records.append(target)
        first, second = copy.deepcopy(records), copy.deepcopy(records)
        report = assign_reviewed_split(first)
        assign_reviewed_split(second)
        self.assertEqual(first, second)
        self.assertEqual(first[1]["split"], "train")
        memberships = {split: {r["group"] for r in first if r["split"] == split} for split in ("train", "dev", "test")}
        self.assertFalse(memberships["train"] & memberships["dev"])
        self.assertFalse(memberships["train"] & memberships["test"])
        self.assertFalse(memberships["dev"] & memberships["test"])
        self.assertTrue(report["all_splits_contain_both_binary_classes"])
        self.assertFalse(report["source_session_independence_confirmed"])

    def test_changed_video_rejected_against_workbook_sha_before_cache_or_decode(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "videos").mkdir()
            (root / "videos/X-1.mov").write_bytes(b"changed-video")
            path = root / "review.xlsx"
            workbook_fixture(path, [["X-1.mov", "нарушение", 2, None, "Подтверждено", "нарушение", 2,
                "Не отдал полную сумму", None, "Human note", None]],
                [["X-1.mov", hashlib.sha256(b"reviewed-original").hexdigest(), 1, 1, None]])
            with patch("cashwatch.reviewed_data.decode_sampled") as decoder:
                with self.assertRaisesRegex(ValueError, "SHA256 differs"):
                    prepare_reviewed(path, root / "videos", root / "output", human_action_notes_confirmed=True)
                decoder.assert_not_called()
            self.assertFalse((root / "output").exists())

    def test_minimal_preparation_keeps_provenance_and_reports_insufficient_holdout_support(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "videos").mkdir()
            for name in ("X-1.mov", "Х-3.mov"):
                (root / "videos" / name).write_bytes(name.encode())
            path = root / "review.xlsx"
            workbook_fixture(path, [
                ["X-1.mov", "нарушение", 2, None, "Подтверждено", "нарушение", 2, "Не отдал полную сумму", None, "Human note", None],
                ["Х-3.mov", "без нарушения", None, None, "Подтверждено", "без нарушения", None, "без нарушения", None, "Human note", None]])
            pixels = np.zeros((6, 224, 224, 3), dtype=np.uint8)
            with patch("cashwatch.reviewed_data.decode_sampled", return_value=(pixels, np.arange(6) / 2, 3.0)):
                manifest = prepare_reviewed(path, root / "videos", root / "output", human_action_notes_confirmed=True)
            self.assertEqual(len(manifest["records"]), 2)
            self.assertFalse(manifest["split_plan"]["all_splits_contain_both_binary_classes"])
            self.assertFalse(manifest["independent_test"])
            self.assertEqual(manifest["group_review"]["status"], "pending")
            self.assertTrue(all(Path(r["cache"]).exists() for r in manifest["records"]))
            self.assertTrue((root / "output/manifest.json").is_file())


if __name__ == "__main__":
    unittest.main()
