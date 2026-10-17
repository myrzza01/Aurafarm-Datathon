"""Training-only labels and preprocessing. Never imported by the predictor."""
from __future__ import annotations

import csv
import copy
import datetime as dt
import hashlib
import json
import random
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .video import decode_sampled


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def canonical_name(value):
    return Path(str(value).strip()).stem.lower().translate(str.maketrans({"х": "x", "–": "-", "—": "-"}))


def parse_label(value):
    value = re.sub(r"\s+", " ", str(value).strip().lower())
    labels = {"нарушение": 1, "без нарушения": 0, "violation": 1, "no_violation": 0, "1": 1, "0": 0}
    if value not in labels:
        raise ValueError(f"Unknown binary label: {value!r}")
    return labels[value]


def parse_timestamps(value):
    if value is None or value == "":
        return []
    if isinstance(value, dt.time):
        return [value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6]
    if isinstance(value, (float, int)):
        return [float(value)]  # Numeric values are seconds, Excel time cells are dt.time.
    result = []
    for part in re.split(r"[;\n]", str(value).strip()):
        part = re.sub(r"\s*(секунда|секунды|секунд|сек|seconds?|s)\.?\s*$", "", part.strip(), flags=re.I)
        if re.fullmatch(r"\d+(?:[.,]\d+)?", part):
            result.append(float(part.replace(",", ".")))
        elif re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?", part):
            total = 0.0
            for number in part.split(":"):
                total = total * 60 + float(number)
            result.append(total)
        else:
            raise ValueError(f"Ambiguous timestamp: {value!r}; use seconds or mm:ss, separated by semicolons")
    return sorted(set(result))


def read_annotations(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if len(header) < 3 or "файл" not in str(header[0]).lower():
            raise ValueError("Expected columns: filename, timestamp, binary label")
        records = {}
        for row_number, row in enumerate(rows, 2):
            if not row or row[0] is None:
                continue
            name = canonical_name(row[0])
            raw_label = re.sub(r"\s+", " ", str(row[2]).strip().lower())
            label = None if raw_label in {"неясно", "uncertain", "unknown"} else parse_label(row[2])
            times = parse_timestamps(row[1])
            if label == 1 and not times:
                raise ValueError(f"Positive row {row_number} needs an event timestamp")
            if label == 0 and times:
                raise ValueError(f"Negative row {row_number} unexpectedly has timestamps")
            if name in records:
                if records[name]["label"] != label:
                    raise ValueError(f"Conflicting labels for {name}")
                records[name]["event_times"] = sorted(set(records[name]["event_times"] + times))
            else:
                records[name] = {"id": name, "label": label, "event_times": times}
                if label is None:
                    records[name]["annotation_status"] = "uncertain_quarantined"
        return list(records.values())
    finally:
        workbook.close()


def grouped_split(records, seed=42, val_fraction=0.25):
    """Video/group split; keeps identical file bytes together even under aliases."""
    groups = sorted({r["group"] for r in records})
    rng = random.Random(seed)
    best = None
    for _ in range(2000):
        shuffled = groups.copy()
        rng.shuffle(shuffled)
        cut = max(1, round(len(groups) * val_fraction))
        val_groups = set(shuffled[:cut])
        train = [r for r in records if r["group"] not in val_groups]
        val = [r for r in records if r["group"] in val_groups]
        if {r["label"] for r in train} != {0, 1} or {r["label"] for r in val} != {0, 1}:
            continue
        score = abs(len(val) / len(records) - val_fraction)
        score += abs(sum(r["label"] for r in val) / len(val) - sum(r["label"] for r in records) / len(records))
        if best is None or score < best[0]:
            best = (score, val_groups)
    if best is None:
        raise ValueError("Need enough independent positive and negative groups for both train and validation")
    for record in records:
        record["split"] = "val" if record["group"] in best[1] else "train"
    return records


def preserve_split(records, previous, seed):
    """Keep existing membership; new groups go to train, aliases inherit their group."""
    if not previous:
        grouped_split(records, seed)
        for record in records:
            record["split"] = "dev" if record["split"] == "val" else record["split"]
        return
    if previous["seed"] != seed:
        raise ValueError("Cannot change seed of an existing split registry")
    old = {r["id"]: r for r in previous["records"]}
    old_groups = {}
    for record in previous["records"]:
        split = "dev" if record["split"] == "val" else record["split"]
        old_groups.setdefault(record["group"], set()).add(split)
    for record in records:
        if record["id"] in old and old[record["id"]]["sha256"] != record["sha256"]:
            raise ValueError("An existing video changed bytes; use a new ID and explicit source group")
    # Group merging is allowed only when all previous assignments agree.
    groups = {}
    for record in records:
        inherited = set(old_groups.get(record["group"], set()))
        if record["id"] in old:
            value = old[record["id"]]["split"]
            inherited.add("dev" if value == "val" else value)
        groups.setdefault(record["group"], set()).update(inherited)
    for record in records:
        inherited = groups[record["group"]]
        if len(inherited) > 1:
            raise ValueError("Related videos cross the locked split. Resolve grouping and start a separately reviewed split")
        record["split"] = next(iter(inherited)) if inherited else "train"
    old_dev = {r["id"] for r in previous["records"] if r["split"] in {"dev", "val", "test"}}
    if old_dev - {r["id"] for r in records}:
        raise ValueError("Locked development/test videos are missing from the updated dataset")


def resolve_groups(records, previous, custom_groups, signatures, related_pairs=()):
    """Union explicit source groups, exact bytes and conservative visual duplicates."""
    parent = {r["id"]: r["id"] for r in records}
    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    def union(a, b):
        parent[find(b)] = find(a)
    old = {r["id"]: r for r in (previous or {}).get("records", [])}
    seen = {}
    for record in records:
        keys = [("sha", record["sha256"])]
        if record["id"] in old:
            keys.append(("old_group", old[record["id"]]["group"]))
        if record["id"] in custom_groups:
            keys.append(("source_group", custom_groups[record["id"]]))
        for key in keys:
            if key in seen:
                union(record["id"], seen[key])
            else:
                seen[key] = record["id"]
    for first, second in related_pairs:
        union(first, second)
    duplicate_pairs = []
    for i, first in enumerate(records):
        for second in records[i+1:]:
            if abs(first["duration_sec"] - second["duration_sec"]) > 0.1:
                continue
            diff = abs(signatures[first["id"]].astype(float) - signatures[second["id"]].astype(float))
            # Whole-sequence near-equality, not scene/background similarity.
            if float(diff.mean()) < 1.5 and float(np.quantile(diff, 0.95)) < 5:
                union(first["id"], second["id"])
                duplicate_pairs.append([first["id"], second["id"]])
    members = {}
    for record in records:
        members.setdefault(find(record["id"]), []).append(record)
    for group_records in members.values():
        prior_groups = sorted({old[r["id"]]["group"] for r in group_records if r["id"] in old})
        explicit = sorted({custom_groups[r["id"]] for r in group_records if r["id"] in custom_groups})
        group = prior_groups[0] if prior_groups else ("source:" + explicit[0] if explicit else min(r["sha256"] for r in group_records))
        for record in group_records:
            record["group"] = group
    by_hash = {}
    for record in records:
        annotation = (record["label"], tuple(record["event_times"]))
        if record["sha256"] in by_hash and by_hash[record["sha256"]] != annotation:
            raise ValueError("Identical video bytes have conflicting annotations")
        by_hash[record["sha256"]] = annotation
    return duplicate_pairs


def audit_sequence_pair(first, second, first_times, second_times, min_frames=8):
    """Find aligned consecutive frames, including temporal crops and overlaps.

    Inputs are locally computed block-averaged RGB frames. Similar background
    alone is insufficient for an automatic merge: a strict match also needs
    matching temporal variation. Ambiguous matches are excluded from new test
    candidates rather than declared independent.
    """
    a, b = np.asarray(first, dtype=np.float32), np.asarray(second, dtype=np.float32)
    if min(len(a), len(b)) < min_frames:
        return None
    distances = np.empty((len(a), len(b)), dtype=np.float32)
    for start in range(0, len(a), 16):
        distances[start:start+16] = np.abs(a[start:start+16, None] - b[None]).mean(axis=(2, 3, 4))
    best = None
    for offset in range(-len(a) + min_frames, len(b) - min_frames + 1):
        ia = np.arange(max(0, -offset), min(len(a), len(b) - offset))
        ib = ia + offset
        good = distances[ia, ib] <= 8.0
        # A shared static background can make pixel distances small for entire
        # unrelated takes. Require aligned changes, not just low absolute MAE.
        da_all, db_all = np.diff(a[ia], axis=0), np.diff(b[ib], axis=0)
        activity = (np.abs(da_all).mean(axis=(1, 2, 3)) + np.abs(db_all).mean(axis=(1, 2, 3))) / 2
        change_error = np.abs(da_all - db_all).mean(axis=(1, 2, 3))
        norm = np.sqrt((da_all ** 2).sum(axis=(1, 2, 3)) * (db_all ** 2).sum(axis=(1, 2, 3)))
        similarity = (da_all * db_all).sum(axis=(1, 2, 3)) / np.maximum(norm, 1e-8)
        aligned_change = (activity >= 0.15) & (similarity >= 0.8) & (change_error <= activity * 0.8)
        near_identical_static = (activity < 0.15) & (distances[ia[1:], ib[1:]] <= 0.25)
        good[1:] &= aligned_change | near_identical_static
        good[0] &= good[1]
        edges = np.flatnonzero(np.diff(np.r_[False, good, False]))
        for start, stop in zip(edges[::2], edges[1::2]):
            if stop - start < min_frames:
                continue
            x, y = ia[start:stop], ib[start:stop]
            delta_time = np.asarray(second_times)[y] - np.asarray(first_times)[x]
            if float(np.ptp(delta_time)) > 0.12:
                continue
            diff = np.abs(a[x] - b[y])
            mean, p95 = float(diff.mean()), float(np.quantile(diff, 0.95))
            if p95 > 24:
                continue
            da, db = np.diff(a[x], axis=0), np.diff(b[y], axis=0)
            motion = float((np.abs(da).mean() + np.abs(db).mean()) / 2)
            motion_error = float(np.abs(da - db).mean())
            cosine = float((da * db).sum() / max(float(np.sqrt((da ** 2).sum() * (db ** 2).sum())), 1e-8))
            strict = mean <= 2.0 and p95 <= 8.0 and motion >= 0.15 and motion_error <= motion * 0.4 and cosine >= 0.95
            candidate = {"status": "confirmed_numeric_overlap" if strict else "possible_numeric_overlap",
                         "matched_frames": len(x), "first_start_sec": float(first_times[x[0]]),
                         "first_end_sec": float(first_times[x[-1]]),
                         "second_start_sec": float(second_times[y[0]]),
                         "second_end_sec": float(second_times[y[-1]]),
                         "offset_sec": float(np.median(delta_time)), "mean_pixel_error": mean,
                         "p95_pixel_error": p95, "temporal_change": motion,
                         "temporal_change_error": motion_error, "temporal_cosine_similarity": cosine}
            rank = (strict, len(x), -mean)
            if best is None or rank > best[0]:
                best = (rank, candidate)
    return best[1] if best else None


def audit_sequences(records, sequences):
    """Metadata-only local audit; never exports images or scene descriptions."""
    matches, exact, compared = [], [], 0
    for i, first in enumerate(records):
        for second in records[i+1:]:
            compared += 1
            if first["sha256"] == second["sha256"]:
                exact.append([first["id"], second["id"]])
                continue
            a, ta = sequences[first["id"]]
            b, tb = sequences[second["id"]]
            match = audit_sequence_pair(a, b, ta, tb)
            if match:
                matches.append({"first_id": first["id"], "second_id": second["id"], **match})
    confirmed = [[m["first_id"], m["second_id"]] for m in matches if m["status"] == "confirmed_numeric_overlap"]
    ambiguous = sorted({value for m in matches if m["status"] == "possible_numeric_overlap"
                        for value in (m["first_id"], m["second_id"])})
    return {"version": 1, "pairs_compared": compared, "exact_byte_pairs": exact,
            "confirmed_overlap_pairs": confirmed, "possible_related_ids": ambiguous,
            "matches": matches, "method": "SHA256 plus consecutive 2-fps block-mean RGB sequence alignment (16x16); min 8 frames",
            "manual_source_review_complete": False,
            "limitations": ["Same recording session or retakes without shared frames cannot be ruled out numerically.",
                            "Spatial crops, speed changes, very short overlaps and strong colour/geometry transformations can escape detection.",
                            "No match means no overlap found under these criteria, not proof of independent provenance."]}


def annotation_revisions(annotations, old_manifest, accept_label_revisions=False):
    old = {r["id"]: r for r in (old_manifest or {}).get("records", []) +
           (old_manifest or {}).get("quarantined_records", [])}
    changes = [{"id": r["id"], "split": old[r["id"]].get("split"),
                "old_label": old[r["id"]]["label"], "new_label": r["label"],
                "old_event_times": old[r["id"]]["event_times"], "new_event_times": r["event_times"]}
               for r in annotations if r["id"] in old and (r["label"], r["event_times"]) !=
               (old[r["id"]]["label"], old[r["id"]]["event_times"])]
    locked = [r for r in changes if r["split"] in {"dev", "val", "test"}]
    if locked and not accept_label_revisions:
        raise ValueError("Locked development/test annotations changed; inspect revisions and explicitly use --accept-label-revisions")
    if any(r["new_label"] is None for r in locked):
        raise ValueError("Locked development/test annotations must retain a definite binary label; uncertain revisions need a separately reviewed split")
    return changes, locked


def group_review_status(records, all_file_ids, old_manifest, annotation_changes,
                        groups_supplied=False, group_review_confirmed=False):
    if group_review_confirmed and not groups_supplied:
        raise ValueError("--confirm-group-review requires --groups-csv supplied by the source owner")
    old = old_manifest or {}
    old_records = old.get("records", []) + old.get("quarantined_records", [])
    old_ids = {r["id"] for r in old_records} | set(old.get("unlabelled_ids", []))
    old_by_id = {r["id"]: r for r in old_records}
    stable_groups = all(r["id"] in old_by_id and r["group"] == old_by_id[r["id"]]["group"] and
                        r["sha256"] == old_by_id[r["id"]]["sha256"] for r in records)
    prior_review = old.get("group_review", {})
    carry = (prior_review.get("status") == "confirmed" and
             prior_review.get("confirmation_method") == "explicit_owner_flag" and
             set(all_file_ids) == old_ids and stable_groups and not annotation_changes)
    confirmed = not annotation_changes and (group_review_confirmed or carry)
    return {"status": "confirmed" if confirmed else "pending",
            "manual_source_review_complete": confirmed,
            "confirmation_method": "explicit_owner_flag" if confirmed else None,
            "reason": ("Source-owner grouping explicitly confirmed" if confirmed else
                       "Annotation revisions require a new review" if annotation_changes else
                       "Source owner has not explicitly confirmed relationships for this complete dataset")}


def prepare(dataset_root, output, labels=None, allow_missing=False, groups_csv=None, seed=42,
            group_review_confirmed=False, accept_label_revisions=False):
    if group_review_confirmed and not groups_csv:
        raise ValueError("--confirm-group-review requires --groups-csv supplied by the source owner")
    root, output = Path(dataset_root).resolve(), Path(output).resolve()
    registry_path = output / "split_registry.json"
    previous_path = registry_path if registry_path.exists() else output / "manifest.json"
    previous = json.loads(previous_path.read_text(encoding="utf-8")) if previous_path.exists() else None
    old_manifest_path = output / "manifest.json"
    old_manifest = json.loads(old_manifest_path.read_text(encoding="utf-8")) if old_manifest_path.exists() else None
    if labels is None:
        candidates = sorted(p for p in root.glob("*.xlsx") if not p.name.startswith("~$"))
        if len(candidates) != 1:
            raise ValueError("Specify --labels when there is not exactly one .xlsx in the dataset root")
        labels = candidates[0]
    labels_digest = file_hash(labels)
    annotations = read_annotations(labels)
    if file_hash(labels) != labels_digest:
        raise ValueError("Labels workbook changed while being read; rerun preparation against a stable workbook")
    annotation_changes, locked_revisions = annotation_revisions(annotations, old_manifest, accept_label_revisions)
    files = {}
    for file in root.rglob("*"):
        if file.is_file() and file.suffix.lower() in {".mov", ".mp4"}:
            key = canonical_name(file.name)
            if key in files:
                raise ValueError(f"Ambiguous normalized filename: {key}")
            files[key] = file
    missing = [r["id"] for r in annotations if r["id"] not in files]
    if missing and not allow_missing:
        raise ValueError(f"Missing {len(missing)} annotated videos. Use --allow-missing for an explicitly partial experiment")
    custom_groups = {}
    if groups_csv:
        with Path(groups_csv).open(encoding="utf-8-sig", newline="") as file:
            custom_groups = {canonical_name(row["filename"]): row["group"] for row in csv.DictReader(file)}
        if set(files) - set(custom_groups):
            raise ValueError("Group CSV must cover every available video")
    output.mkdir(parents=True, exist_ok=True)
    records, hash_groups, signatures, sequences = [], {}, {}, {}
    for item in annotations:
        if item["id"] not in files:
            continue
        path = files[item["id"]]
        digest = file_hash(path)
        group = custom_groups.get(item["id"], digest)
        if digest in hash_groups and hash_groups[digest] != group:
            raise ValueError("Identical videos assigned to different groups")
        hash_groups[digest] = group
        cache = output / "cache" / f"{digest}.npz"
        cache.parent.mkdir(exist_ok=True)
        if cache.exists():
            with np.load(cache, allow_pickle=False) as saved:
                images = saved["images"]
                times = saved["times"]
                duration = float(saved["duration"])
                if int(saved["size"]) != 224 or float(saved["fps"]) != 2:
                    raise ValueError("Cache sampling settings mismatch")
        else:
            images, times, duration = decode_sampled(path)
            np.savez_compressed(cache, images=images, times=times, duration=duration, fps=2.0, size=224)
        if any(not np.isfinite(t) or t < 0 or t > duration for t in item["event_times"]):
            raise ValueError(f"Timestamp outside video duration for {item['id']}")
        records.append({**item, "path": str(path), "sha256": digest, "group": group,
                        "cache": str(cache), "duration_sec": duration, "sampled_frames": len(times)})
        indices = np.linspace(0, len(images)-1, 16).round().astype(int)
        signatures[item["id"]] = images[indices, ::7, ::7]
        # Averaging preserves low-frequency structure while avoiding pixel
        # aliasing. Signatures stay in local memory and are not exported.
        sequences[item["id"]] = (images.reshape(len(images), 16, 14, 16, 14, 3).mean(axis=(2, 4)).astype(np.float32), times)
        print(f"Prepared {len(records)}/{len(annotations) - len(missing)} videos", flush=True)
    if not records:
        raise ValueError("No annotated videos found")
    audit = audit_sequences(records, sequences)
    duplicate_pairs = resolve_groups(records, previous, custom_groups, signatures,
                                     related_pairs=audit["confirmed_overlap_pairs"])
    quarantined = [r for r in records if r["label"] is None]
    records = [r for r in records if r["label"] is not None]
    if not records:
        raise ValueError("No videos with unambiguous binary labels")
    for record in quarantined:
        record["split"] = "quarantine"
    preserve_split(records, previous, seed)
    review = group_review_status(records + quarantined, files, old_manifest, annotation_changes,
                                 bool(groups_csv), group_review_confirmed)
    reservation = copy.deepcopy((previous or {}).get("test_reservation"))
    revisions = copy.deepcopy((old_manifest or {}).get("label_revision_history", []))
    prior_revision_path = None
    if locked_revisions:
        prior_revision_path = output / "revisions" / f"manifest_{file_hash(old_manifest_path)[:16]}.json"
        revisions.append({"previous_manifest": str(prior_revision_path),
                          "previous_labels_sha256": old_manifest.get("labels_sha256"),
                          "new_labels_sha256": labels_digest, "accepted_explicitly": True,
                          "changes": locked_revisions})
        if reservation:
            reservation["qualification_status"] = "invalidated_by_label_revision"
            reservation["label_revisions_pending_review"] = True
    if reservation:
        qualified = review["status"] == "confirmed" and not reservation.get("label_revisions_pending_review", False)
        reservation["status"] = "provenance_confirmed" if qualified else "provenance_pending"
        reservation["source_owner_review_complete"] = review["manual_source_review_complete"]
    manifest = {"version": 3, "seed": seed, "fps": 2.0, "image_size": 224, "clip_frames": 8,
                "positive_radius_sec": 1.0, "records": records, "missing_ids": missing,
                "available_video_count": len(files), "quarantined_records": quarantined,
                "unlabelled_ids": sorted(set(files) - {r["id"] for r in annotations}),
                "groups_supplied": bool(groups_csv) or (bool((previous or {}).get("groups_supplied")) and
                    {r["id"] for r in records} <= {r["id"] for r in (previous or {}).get("records", [])}),
                "duplicate_pairs": duplicate_pairs, "labels_sha256": labels_digest,
                "annotation_changes_since_previous_manifest": annotation_changes,
                "annotation_provenance": (old_manifest or {}).get("annotation_provenance", []),
                "automatic_group_audit": audit, "manual_group_review_complete": review["manual_source_review_complete"],
                "group_review": review, "label_revision_history": revisions,
                "test_reservation": reservation,
                "note": "Locked development and reserved test membership. New groups enter train unless explicitly reserved before fitting. Related takes still need source-owner review."}
    if file_hash(labels) != labels_digest:
        raise ValueError("Labels workbook changed during preparation; manifest was not updated. Rerun against a stable workbook")
    if previous_path.exists() and previous_path.name == "manifest.json":
        write_json(output / "manifest_before_v2.json", previous)
    if prior_revision_path:
        write_json(prior_revision_path, old_manifest)
    write_json(output / "manifest.json", manifest)
    write_json(output / "group_audit.json", audit)
    if reservation:
        write_json(output / "test_reservation.json", reservation)
    # Retain historical entries so temporary removal does not discard known groups.
    historical = {r["id"]: r for r in (previous or {}).get("records", [])}
    historical.update({r["id"]: {k: r[k] for k in ("id", "sha256", "group", "split")} for r in records})
    write_json(registry_path, {"seed": seed, "groups_supplied": manifest["groups_supplied"],
                              "test_reservation": manifest["test_reservation"], "records": list(historical.values())})
    print(json.dumps({"videos": len(records), "positive": sum(r["label"] for r in records),
                      "negative": sum(1-r["label"] for r in records), "missing": len(missing),
                      "train": sum(r["split"] == "train" for r in records),
                      "dev": sum(r["split"] == "dev" for r in records),
                      "test": sum(r["split"] == "test" for r in records),
                      "quarantined": len(quarantined),
                      "duplicate_pairs": len(duplicate_pairs)}), flush=True)
    return manifest
