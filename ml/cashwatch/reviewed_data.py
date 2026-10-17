"""Import reviewed human annotations without using spreadsheet text as features.

Current review cells override the static ML-export sheet. Historical labels remain
explicitly weak unless their review status is confirmed. Action-note confirmation
is a separate, caller-supplied provenance fact; it never confirms source sessions.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import copy
import json
from pathlib import Path
import random
import re

import numpy as np
from openpyxl import load_workbook

from .data import canonical_name, file_hash, parse_label, parse_timestamps, write_json
from .video import decode_sampled


ACTION_TYPES = {
    "cash_not_fully_delivered": "Неполная выдача суммы",
    "customer_identity_not_presented": "Документы клиента не представлены",
    "accepted_cash_left_exposed": "Деньги оставлены вне защищённого хранения",
    "equipment_left_on": "Оборудование оставлено включённым",
    "cash_counting_omitted": "Пересчёт банкнот пропущен",
    "collector_documents_not_presented": "Документы передающего денежный мешок не представлены",
    "simultaneous_customer_service": "Одновременное обслуживание клиентов",
    "safe_left_open": "Сейф оставлен открытым",
    "safe_key_left_in_lock": "Ключ оставлен в замке сейфа",
    "unsecured_storage_on_departure": "Уход с рабочего места с незапертым сейфом",
    "workstation_information_left_open": "Информация оставлена открытой",
    "cash_office_door_left_open": "Дверь кассового кабинета оставлена открытой",
    "customer_departed_before_completion": "Клиент ушёл до завершения операции",
    "personal_items_on_cash_desk": "Личные вещи на кассовом столе",
    "rejected_banknote_recombined_without_check": "Отклонённая банкнота возвращена без проверки",
    "customer_cash_combined": "Банкноты разных клиентов объединены",
}


def canonical_action_type(note):
    """Canonicalize explicit human action wording, never video/file-name features."""
    text = re.sub(r"\s+", " ", str(note or "").strip().lower().replace("ё", "е"))
    patterns = [
        ("cash_not_fully_delivered", r"не отдал полную сумму"),
        ("customer_identity_not_presented", r"клиент не показал документ"),
        ("equipment_left_on", r"не выключила оборудование"),
        ("cash_counting_omitted", r"не посчитала деньги"),
        ("collector_documents_not_presented", r"коллега.*не показывая документ"),
        ("simultaneous_customer_service", r"обслуживает два клиента"),
        ("safe_left_open", r"оставление сейфа открытым"),
        ("safe_key_left_in_lock", r"оставление ключа в замке сейфа"),
        ("unsecured_storage_on_departure", r"оставляет рабочее место"),
        ("workstation_information_left_open", r"оставила информацию"),
        ("cash_office_door_left_open", r"оставила дверь открытой"),
        ("customer_departed_before_completion", r"не досидел до завершения"),
        ("personal_items_on_cash_desk", r"поставила вещи"),
        ("rejected_banknote_recombined_without_check", r"не убрала не\s*пригодную купюру"),
        ("customer_cash_combined", r"объединение банкнот разных клиентов"),
        ("accepted_cash_left_exposed", r"оставила деньги|оставила денежный мешок|не положила их в сейф"),
    ]
    return next((key for key, pattern in patterns if re.search(pattern, text)), None)


def _optional_label(value):
    if value is None or str(value).strip().lower() in {"", "неясно", "unknown", "uncertain"}:
        return None
    return parse_label(value)


def _header_rows(sheet, required):
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        headers = {str(value).strip(): column for column, value in enumerate(row) if value is not None}
        if set(required) <= set(headers):
            return index, headers
    raise ValueError(f"Required annotation columns missing from {sheet.title}")


def _windows_path(path):
    value = str(path)
    match = re.match(r"^/mnt/([a-z])/(.*)$", value)
    return f"{match[1].upper()}:/{match[2]}" if match else value


def read_reviewed_annotations(path, *, allow_weak_original=False, human_action_notes_confirmed=False):
    """Return row annotations plus workbook provenance; never alter the workbook.

    H/J may be confirmed by an explicit user instruction after workbook creation.
    Callers must pass that fact explicitly; old readiness flags remain preserved.
    """
    path = Path(path).resolve()
    digest = file_hash(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        review = workbook["Проверка видео"]
        header_row, headers = _header_rows(review, ["Файл", "Статус проверки", "Метка после проверки", "Тип события", "Комментарий"])
        original = {}
        if "Исходная разметка" in workbook.sheetnames:
            for row_index, row in enumerate(workbook["Исходная разметка"].iter_rows(values_only=True), 1):
                if row_index == 1 or not row or row[0] is None:
                    continue
                original[canonical_name(row[0])] = {"label": _optional_label(row[2]), "times": parse_timestamps(row[1]),
                    "source": {"sheet": "Исходная разметка", "row": row_index, "label_cell": f"C{row_index}", "time_cell": f"B{row_index}"}}
        snapshots = {}
        if "Для ML" in workbook.sheetnames:
            sheet = workbook["Для ML"]
            first, columns = _header_rows(sheet, ["video_file", "sha256", "classification_ready"])
            for row_index, row in enumerate(sheet.iter_rows(min_row=first + 1, values_only=True), first + 1):
                if row[columns["video_file"]] is None:
                    continue
                snapshot = {key: row[column] for key, column in columns.items()}
                snapshot["source_row"] = row_index
                names = [snapshot["video_file"]] + re.split(r"[;,\n]", str(snapshot.get("duplicate_files") or ""))
                for name in names:
                    if name.strip():
                        snapshots[canonical_name(name)] = snapshot
        records, issues = [], []
        for row_index, row in enumerate(review.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            def cell(name):
                column = headers.get(name)
                return row[column] if column is not None and column < len(row) else None
            filename = cell("Файл")
            if filename is None:
                continue
            identifier = canonical_name(filename)
            prior = original.get(identifier, {"label": _optional_label(cell("Исходная метка")),
                "times": parse_timestamps(cell("Исходный таймкод")), "source": {"sheet": review.title, "row": row_index}})
            reviewed_label = _optional_label(cell("Метка после проверки"))
            reviewed_times = parse_timestamps(cell("Таймкод после проверки, с"))
            status = str(cell("Статус проверки") or "").strip()
            confirmed = status.casefold() == "подтверждено" and reviewed_label is not None
            label = reviewed_label if confirmed else prior["label"] if allow_weak_original else None
            times = reviewed_times if confirmed else prior["times"] if allow_weak_original else []
            source = "human_confirmed" if confirmed else "source_original_weak" if label is not None else "unconfirmed_excluded"
            note, description = str(cell("Тип события") or "").strip(), str(cell("Комментарий") or "").strip()
            action_id = canonical_action_type(note) if label == 1 else None
            explicit_time = re.match(r"^\s*(\d+(?:[.,]\d+)?)\s*(?:с|сек(?:унда|унды|унд)?)\s*[—–:-]", note, flags=re.I)
            explicit_action_time = bool(human_action_notes_confirmed and explicit_time)
            if explicit_action_time and label == 1:
                times = [float(explicit_time[1].replace(",", "."))]
            if label == 0:
                times = []
            provenance = {"workbook_sha256": digest, "sheet": review.title, "row": row_index,
                "label_cell": f"F{row_index}", "time_cell": f"G{row_index}", "action_cell": f"H{row_index}",
                "description_cell": f"J{row_index}", "reviewer": cell("Проверяющий"), "review_status": status}
            action_verified = bool(human_action_notes_confirmed and note)
            legal_urls = sorted(set(re.findall(r"https?://[^\s<>\[\]()]+", note + " " + description)))
            events = [{"timecode_sec": timestamp, "type_id": action_id, "type_verified": bool(action_verified and action_id),
                "event_type_note": note, "observed_description": description,
                "description_source": "user_confirmed_human_annotation" if action_verified else "unconfirmed_human_note",
                "legal_sources": [{"url": url, "status": "provided_unverified"} for url in legal_urls],
                "time_verified": bool(confirmed or explicit_action_time),
                "time_source": "human_action_note_explicit_point" if explicit_action_time else source,
                "annotation_provenance": copy.deepcopy(provenance)} for timestamp in times] if label == 1 else []
            if confirmed and label == 1 and not times:
                issues.append({"id": identifier, "reason": "confirmed_positive_missing_time"})
            if action_verified and label == 1 and action_id is None:
                issues.append({"id": identifier, "reason": "action_note_has_no_canonical_type"})
            record = {"id": identifier, "filename": str(filename), "label": label, "event_times": times, "events": events,
                "label_source": source, "label_verified": confirmed, "review_status": status,
                "annotation_confidence": source, "event_time_verified": bool(confirmed or explicit_action_time),
                "event_type_note": note, "observed_description": description,
                "type_id": action_id, "type_verified": bool(action_verified and action_id),
                "human_action_verified": action_verified, "legal_sources": legal_urls,
                "annotation_provenance": [provenance], "original_annotation": prior,
                "reviewed_annotation": {"label": reviewed_label, "event_times": reviewed_times},
                "source_snapshot": snapshots.get(identifier), "data_note": cell("Примечание к данным"),
                "action_confirmation_source": "explicit_user_instruction_in_current_conversation" if human_action_notes_confirmed else None}
            records.append(record)
        if file_hash(path) != digest:
            raise ValueError("Workbook changed while importing annotations")
        return {"workbook": str(path), "workbook_sha256": digest, "records": records, "issues": issues,
            "authoritative_sheet": review.title, "review_status_counts": dict(Counter(r["review_status"] for r in records)),
            "allow_weak_original": allow_weak_original, "human_action_notes_confirmed": human_action_notes_confirmed,
            "source_priority": "Confirmed current review F/G, then explicitly allowed weak original labels; static ML snapshot is metadata only."}
    finally:
        workbook.close()


def reconcile_exact_duplicates(records):
    """One sample per SHA; retain compatible distinct action events and every source."""
    groups = defaultdict(list)
    for record in records:
        groups[record["sha256"]].append(record)
    selected, quarantined, audit = [], [], []
    for digest, members in sorted(groups.items()):
        reason = None
        verified = {r["label"] for r in members if r["label_verified"] and r["label"] is not None}
        labels = verified or {r["label"] for r in members if r["label"] is not None}
        if len(labels) > 1:
            reason = "conflicting_confirmed_labels" if verified else "unresolved_original_label_conflict"
        label = next(iter(labels)) if len(labels) == 1 else None
        ranked = sorted(members, key=lambda r: (not r["label_verified"],
            int(re.search(r"\d+", r["id"])[0]) if re.search(r"\d+", r["id"]) else float("inf"), r["id"]))
        chosen = copy.deepcopy(ranked[0])
        chosen.update({"group": digest, "aliases": [r["id"] for r in ranked],
            "alias_paths": [r["path"] for r in ranked], "label": label,
            "annotation_provenance": [source for r in ranked for source in r["annotation_provenance"]],
            "alias_annotations": [{key: r.get(key) for key in ["id", "label", "label_source", "label_verified", "review_status", "event_times", "events", "original_annotation"]} for r in ranked]})
        candidates = [r for r in ranked if r["label"] == label]
        if label == 1 and not reason:
            typed = [r for r in candidates if r["type_verified"] and r["events"]]
            if len(typed) == len(candidates) and len({r["type_id"] for r in typed}) > 1:
                by_type = defaultdict(list)
                for record in typed:
                    by_type[record["type_id"]].append(record)
                if any(len({tuple(r["event_times"]) for r in values}) > 1 for values in by_type.values()):
                    reason = "same_confirmed_action_has_unresolved_times"
                else:
                    chosen["events"] = [copy.deepcopy(event) for values in by_type.values() for event in values[0]["events"]]
            else:
                strong = [r for r in candidates if r["label_verified"]]
                used = strong or candidates
                if len({tuple(r["event_times"]) for r in used}) > 1:
                    reason = "unresolved_duplicate_event_time_conflict"
                else:
                    chosen["events"] = copy.deepcopy(used[0]["events"]) if used else []
            chosen["events"] = sorted(chosen["events"], key=lambda e: (e["timecode_sec"], e["type_id"] or ""))
            chosen["event_times"] = sorted({event["timecode_sec"] for event in chosen["events"]})
            if not chosen["event_times"]:
                reason = reason or "positive_without_event_time"
        elif label == 0:
            chosen["events"], chosen["event_times"] = [], []
        elif not reason:
            reason = "no_authorized_binary_label"
        if reason:
            chosen.update({"split": "quarantine", "exclusion_reason": reason})
            quarantined.append(chosen)
        else:
            chosen["exclusion_reason"] = None
            selected.append(chosen)
        if len(members) > 1 or reason:
            audit.append({"sha256": digest, "ids": chosen["aliases"], "status": "quarantined" if reason else "resolved",
                "reason": reason or ("distinct_user_confirmed_actions_merged" if len({e["type_id"] for e in chosen["events"]}) > 1 else "confirmed_review_precedence_or_consistent_annotations"),
                "selected_label": label, "event_times": chosen["event_times"],
                "type_ids": [e["type_id"] for e in chosen["events"]], "sample_weight_count": 1})
    return selected, quarantined, audit


def assign_reviewed_split(records, seed=42, holdout_fraction=0.15):
    """Stratify binary groups; keep singleton action categories in training.

    These are new diagnostic partitions, not an assertion of source-session or
    historical-training independence. Every known action remains represented in
    training; the support report exposes categories with too few groups.
    """
    rng = random.Random(seed)
    groups = defaultdict(list)
    for record in records:
        groups[record["group"]].append(record)
    classes = {key: {r["label"] for r in values} for key, values in groups.items()}
    if any(len(value) != 1 for value in classes.values()):
        raise ValueError("A reviewed split group contains conflicting video labels")
    type_groups = defaultdict(set)
    for record in records:
        for event in record["events"]:
            if event["type_verified"]:
                type_groups[event["type_id"]].add(record["group"])
    protected = set().union(*(values for values in type_groups.values() if len(values) < 3)) if type_groups else set()
    membership = {key: "train" for key in groups}
    available = {label: [key for key, values in classes.items() if values == {label} and key not in protected] for label in (0, 1)}
    targets = {label: max(1, round(sum(values == {label} for values in classes.values()) * holdout_fraction)) for label in (0, 1)}
    for values in available.values():
        rng.shuffle(values)
    for split in ("dev", "test"):
        for label in (0, 1):
            for _ in range(targets[label]):
                candidates = []
                for order, group in enumerate(available[label]):
                    if membership[group] != "train":
                        continue
                    train_class = [key for key, assigned in membership.items() if assigned == "train" and classes[key] == {label}]
                    if len(train_class) <= 1:
                        continue
                    attached = [values for values in type_groups.values() if group in values]
                    # Reserve one train and one future test group while choosing
                    # dev. Otherwise greedy dev placement can consume every
                    # available example of a three-group action category.
                    minimum_remaining = 2 if split == "dev" else 1
                    if any(sum(membership[key] == "train" for key in values) <= minimum_remaining for values in attached):
                        continue
                    already_present = sum(sum(membership[key] == split for key in values) for values in attached)
                    scarcity = min((len(values) for values in attached), default=float("inf"))
                    candidates.append(((already_present, scarcity, order), group))
                if not candidates:
                    break
                membership[min(candidates)[1]] = split
    for record in records:
        record["split"] = membership[record["group"]]
    coverage = {split: dict(Counter(r["label"] for r in records if r["split"] == split)) for split in ("train", "dev", "test")}
    return {"method": "seeded binary stratification with rare-action train protection", "seed": seed,
        "protected_rare_type_groups": sorted(protected), "binary_counts": coverage,
        "all_splits_contain_both_binary_classes": all(set(values) == {0, 1} for values in coverage.values()),
        "group_disjoint": True, "source_session_independence_confirmed": False,
        "prior_training_overlap": "unknown; prior model manifests unavailable"}


def prepare_reviewed(workbook_path, video_dir, output, *, seed=42, allow_weak_original=False,
                     human_action_notes_confirmed=False, fps=2.0, image_size=224, path_style="native"):
    """Create a local manifest/cache with exact-SHA deduplication and source trails."""
    if path_style not in {"native", "windows"}:
        raise ValueError("path_style must be native or windows")
    output, video_dir = Path(output).resolve(), Path(video_dir).resolve()
    if (output / "manifest.json").exists():
        raise FileExistsError("Choose a fresh reviewed-data output to preserve the split and provenance")
    imported = read_reviewed_annotations(workbook_path, allow_weak_original=allow_weak_original,
        human_action_notes_confirmed=human_action_notes_confirmed)
    paths = {}
    for path in video_dir.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".mov", ".mp4"}:
            identifier = canonical_name(path)
            if identifier in paths:
                raise ValueError(f"Ambiguous canonical video name: {identifier}")
            paths[identifier] = path
    missing = sorted({r["id"] for r in imported["records"]} - paths.keys())
    if missing:
        raise FileNotFoundError(f"Missing {len(missing)} reviewed videos")
    for record in imported["records"]:
        path = paths[record["id"]]
        digest = file_hash(path)
        record.update({"path": str(path), "path_windows": _windows_path(path), "sha256": digest})
        snapshot = record["source_snapshot"]
        if snapshot and snapshot.get("sha256") and snapshot["sha256"] != digest:
            raise ValueError(f"Source video SHA256 differs from reviewed workbook: {record['id']}")
    records, quarantined, resolutions = reconcile_exact_duplicates(imported["records"])
    split_plan = assign_reviewed_split(records, seed)
    output.mkdir(parents=True, exist_ok=True)
    cache_dir = output / "cache"
    cache_dir.mkdir(exist_ok=True)
    for index, record in enumerate(records, 1):
        images, times, duration = decode_sampled(record["path"], fps=fps, size=image_size)
        if any(not np.isfinite(t) or t < 0 or t > duration for t in record["event_times"]):
            raise ValueError(f"Reviewed event lies outside decoded video: {record['id']}")
        cache = cache_dir / f"{record['sha256']}.npz"
        np.savez_compressed(cache, images=images, times=times, duration=duration, fps=fps, size=image_size)
        record.update({"cache": str(cache), "cache_windows": _windows_path(cache), "duration_sec": duration,
            "cache_sha256": file_hash(cache), "sampled_frames": len(times), "classification_ready": True,
            "semantic_ready": any(event["type_verified"] for event in record["events"])})
        print(json.dumps({"prepared_reviewed_videos": index, "total": len(records)}), flush=True)
    support = {}
    for type_id, label in ACTION_TYPES.items():
        matching = [r for r in records if any(e["type_id"] == type_id and e["type_verified"] for e in r["events"])]
        counts = {split: len({r["group"] for r in matching if r["split"] == split}) for split in ("train", "dev", "test")}
        support[type_id] = {"label": label, "groups": len({r["group"] for r in matching}), "by_split": counts,
            "minimum_recommended_train_groups": 2, "minimum_train_support_met": counts["train"] >= 2,
            "all_splits_have_support": all(counts.values()), "generalization_established": False}
    if file_hash(workbook_path) != imported["workbook_sha256"]:
        raise ValueError("Workbook changed during reviewed preparation")
    if any(file_hash(r["path"]) != r["sha256"] for r in imported["records"]):
        raise ValueError("Source video changed during reviewed preparation")
    manifest = {"version": 4, "seed": seed, "fps": fps, "image_size": image_size, "clip_frames": 8,
        "positive_radius_sec": 1.0, "records": records, "quarantined_records": quarantined,
        "missing_ids": [], "unlabelled_ids": sorted(paths.keys() - {r["id"] for r in imported["records"]}),
        "available_video_count": len(paths), "labels_sha256": imported["workbook_sha256"],
        "annotation_source": {key: value for key, value in imported.items() if key != "records"},
        "duplicate_resolutions": resolutions, "split_plan": split_plan,
        "group_review": {"status": "pending", "manual_source_review_complete": False,
            "reason": "Exact SHA groups verified; recording-session relationships and prior-training overlap are not established."},
        "manual_group_review_complete": False, "independent_test": False,
        "test_reservation": {"status": "diagnostic_holdout", "reserved_before_this_training": True,
            "independent_of_prior_training": False},
        "event_taxonomy": {"status": "human_action_notes_confirmed" if human_action_notes_confirmed else "unverified_notes",
            "type_ids": [key for key in ACTION_TYPES if support[key]["groups"]], "types": ACTION_TYPES, "support": support},
        "limitations": ["Original labels outside confirmed review remain weak, even when action notes are user-confirmed.",
            "Descriptions are human annotations for target/provenance only; never model input features.",
            "A point annotation does not establish event onset/end or a legal conclusion.",
            "Held-out partitions are diagnostic; source-session independence and historical-training overlap remain unknown.",
            "Singleton semantic categories are kept in train and have no generalization evidence."],
        "source_files_verified_unchanged": True}
    if path_style == "windows":
        for record in records:
            record["path"], record["cache"] = record["path_windows"], record["cache_windows"]
            record["alias_paths"] = [_windows_path(path) for path in record["alias_paths"]]
    write_json(output / "manifest.json", manifest)
    write_json(output / "split_registry.json", {"seed": seed, "records": [{key: r[key] for key in ("id", "sha256", "group", "split")} for r in records]})
    write_json(output / "import_report.json", {"source_workbook_sha256": imported["workbook_sha256"],
        "input_annotation_rows": len(imported["records"]), "unique_usable_videos": len(records), "quarantined_groups": len(quarantined),
        "source_binary_labels": dict(Counter(r["label_source"] for r in records)),
        "review_status_counts": imported["review_status_counts"], "duplicate_resolutions": resolutions,
        "split_plan": split_plan, "type_support": support, "issues": imported["issues"],
        "source_files_verified_unchanged": True})
    return manifest
